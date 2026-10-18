
import csv
import json
from pathlib import Path

import pymupdf
import pytesseract
from PIL import Image, ImageOps, ImageEnhance, ImageFilter


# =========================================================
# CONFIGURATION
# =========================================================

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".pptx",
    ".txt", ".csv", ".json", ".py",
    ".jpg", ".jpeg", ".png", ".bmp",
    ".tif", ".tiff",
}


# =========================================================
# IMAGE PREPROCESSING
# =========================================================

def preprocess_image(image):
    """
    Prepare an image for OCR by correcting orientation,
    enlarging small text, converting to grayscale,
    and improving contrast.
    """
    image = ImageOps.exif_transpose(image)
    image = image.convert("RGB")

    # Enlarge images when the text may be too small.
    if image.width < 1600:
        scale = min(3, 1600 / image.width)
        image = image.resize(
            (
                int(image.width * scale),
                int(image.height * scale),
            ),
            Image.Resampling.LANCZOS,
        )

    # Convert to grayscale.
    image = ImageOps.grayscale(image)

    # Improve contrast and lightly sharpen text.
    image = ImageEnhance.Contrast(image).enhance(1.8)
    image = image.filter(ImageFilter.SHARPEN)

    # Thresholding can improve text on clean receipt backgrounds.
    image = image.point(
        lambda pixel: 255 if pixel > 170 else 0
    )

    return image


# =========================================================
# OCR FOR IMAGES
# =========================================================

def ocr_image(image):
    """Extract text from an image using Tesseract OCR."""
    try:
        prepared = preprocess_image(image)

        # Preserve reading order and layout reasonably well.
        text = pytesseract.image_to_string(
            prepared,
            lang="eng",
            config="--oem 3 --psm 6",
        )

        return text.strip()

    except Exception as error:
        print(f"OCR error: {error}")
        return ""


# =========================================================
# OCR FOR PDF PAGES
# =========================================================

def extract_pdf_text(file_path):
    """
    Extract embedded text first.
    Use OCR when a PDF page has little embedded text.
    """
    extracted_pages = []

    try:
        with pymupdf.open(str(file_path)) as document:
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text").strip()

                if len(page_text) < 30:
                    try:
                        pixmap = page.get_pixmap(
                            matrix=pymupdf.Matrix(3, 3),
                            alpha=False,
                        )

                        image = Image.frombytes(
                            "RGB",
                            (pixmap.width, pixmap.height),
                            pixmap.samples,
                        )

                        ocr_text = ocr_image(image)

                        if len(ocr_text) > len(page_text):
                            page_text = ocr_text

                    except Exception as error:
                        print(
                            f"OCR failed on PDF page {page_number}: "
                            f"{error}"
                        )

                extracted_pages.append(
                    f"--- Page {page_number} ---\n{page_text}"
                )

    except Exception as error:
        print(f"PDF extraction failed: {error}")
        return ""

    return "\n\n".join(extracted_pages).strip()


# =========================================================
# GENERAL TEXT EXTRACTION
# =========================================================

def extract_text(file_path):
    """
    Extract text from supported document formats.
    Scanned PDFs and images use OCR.
    """
    path = Path(file_path)

    if not path.is_file():
        print(f"File not found: {path}")
        return ""

    extension = path.suffix.lower()

    try:
        if extension == ".pdf":
            return extract_pdf_text(path)

        if extension in {
            ".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"
        }:
            with Image.open(path) as image:
                return ocr_image(image)

        if extension == ".docx":
            from docx import Document

            document = Document(str(path))
            parts = []

            for paragraph in document.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text)

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(parts).strip()

        if extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                str(path),
                read_only=True,
                data_only=True,
            )

            parts = []

            try:
                for sheet in workbook.worksheets:
                    parts.append(f"Sheet: {sheet.title}")

                    for row in sheet.iter_rows(values_only=True):
                        values = [
                            str(value)
                            for value in row
                            if value is not None
                        ]

                        if values:
                            parts.append(" | ".join(values))
            finally:
                workbook.close()

            return "\n".join(parts).strip()

        if extension == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(path))
            parts = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1,
            ):
                parts.append(f"Slide {slide_number}")

                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        if shape.text.strip():
                            parts.append(shape.text)

                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            parts.append(
                                " | ".join(
                                    cell.text for cell in row.cells
                                )
                            )

            return "\n".join(parts).strip()

        if extension in {".txt", ".py", ".json", ".csv"}:
            text = path.read_text(
                encoding="utf-8-sig",
                errors="replace",
            )

            if extension == ".json":
                try:
                    return json.dumps(
                        json.loads(text),
                        indent=2,
                        ensure_ascii=False,
                    )
                except json.JSONDecodeError:
                    return text

            if extension == ".csv":
                try:
                    rows = []

                    with path.open(
                        "r",
                        encoding="utf-8-sig",
                        errors="replace",
                        newline="",
                    ) as file:
                        for row in csv.reader(file):
                            rows.append(" | ".join(row))

                    return "\n".join(rows)
                except Exception:
                    return text

            return text

        return ""

    except Exception as error:
        print(f"Text extraction failed for {path}: {error}")
        return ""


# =========================================================
# DIRECTORY SCANNER
# =========================================================

def scan_directory(folder_path):
    """
    Recursively scan a folder and return file metadata.
    This function does not move, rename, or delete files.
    """
    root = Path(folder_path).expanduser()

    if not root.is_dir():
        raise NotADirectoryError(
            f"Folder does not exist: {root}"
        )

    results = []

    for path in root.rglob("*"):
        try:
            if not path.is_file():
                continue

            if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue

            stat = path.stat()

            results.append({
                "filename": path.name,
                "extension": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified_time": stat.st_mtime,
                "full_path": str(path.resolve()),
            })

        except OSError as error:
            print(f"Could not inspect {path}: {error}")

    return results
